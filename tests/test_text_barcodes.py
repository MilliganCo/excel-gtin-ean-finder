import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook

root = Path(__file__).parents[1]
script = root / 'analyze_barcodes.py'
if not script.exists():
    script = root / 'analyze_barcodes.py.py'
spec = importlib.util.spec_from_file_location('analyze_barcodes', script)
analyzer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(analyzer)


class BarcodeTests(unittest.TestCase):
    def test_text_zeros_unique_counts_numeric_columns_and_original_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'synthetic.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['Штрихкод', 'Номенклатура', 'Цена'])
            sheet.append(['0012345678905', 'Example A', 12.5])
            sheet.append(['12345678905', 'Example A', 20])
            sheet.append(['0000', 'Example B', 0])
            workbook.save(source)
            before = hashlib.sha256(source.read_bytes()).digest()
            output, bad_count = analyzer.analyze_file(str(source))
            self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), before)
            self.assertEqual(bad_count, 1)
            result = load_workbook(output, read_only=True, data_only=True)
            rows = list(result['Все данные с анализом'].iter_rows(values_only=True))
            self.assertEqual([row[0] for row in rows[1:]], ['0012345678905', '12345678905', '0000'])
            self.assertEqual([row[3] for row in rows[1:]], [2, 2, 1])
            self.assertEqual(rows[1][2], 12.5)
            self.assertIsInstance(rows[1][2], (int, float))
            self.assertEqual(list(result['Только 1 штрихкод'].values)[1][0], 'Example B')
            result.close()


if __name__ == '__main__':
    unittest.main()
